"""File validation and text extraction. Never logs resume text."""
import io
import re
import zipfile

from constants import MAX_FILE_BYTES, MAX_TEXT_CHARS, MIN_TEXT_CHARS
from errors import ApiError

ALLOWED_EXT = {".pdf": "pdf", ".docx": "docx", ".doc": "doc"}
OLE_MAGIC = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"


def sanitize_filename(name: str) -> str:
    name = (name or "resume").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^A-Za-z0-9._-]", "_", name).lstrip(".")
    name = re.sub(r"\.{2,}", ".", name)
    return (name or "resume")[:120]


def detect_kind(filename: str, data: bytes) -> str:
    """Validate size, extension and magic bytes. Returns 'pdf' | 'docx' | 'doc'."""
    if len(data) == 0:
        raise ApiError(422, "empty_file", "The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ApiError(413, "file_too_large", "File is larger than 4 MB. Please compress it.")
    ext = "." + (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    kind = ALLOWED_EXT.get(ext)
    if not kind:
        raise ApiError(415, "unsupported_type", "Only PDF, DOCX or DOC files are supported.")
    if kind == "pdf" and not data[:1024].lstrip().startswith(b"%PDF"):
        raise ApiError(415, "bad_signature", "This file is not a valid PDF.")
    if kind == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise ValueError
        except Exception:
            raise ApiError(415, "bad_signature", "This file is not a valid DOCX.")
    if kind == "doc" and not data.startswith(OLE_MAGIC):
        raise ApiError(415, "bad_signature", "This file is not a valid DOC.")
    return kind


def _normalise(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise ApiError(422, "unreadable", "This PDF could not be read.")
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise ApiError(422, "password_protected", "This PDF is password protected.")
    parts = []
    try:
        for page in reader.pages:
            parts.append(page.extract_text() or "")
    except Exception:
        parts = []
    text = "\n".join(parts)
    if len(text.strip()) < MIN_TEXT_CHARS:
        try:  # fall back to pdfplumber for odd encodings
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                alt = "\n".join((p.extract_text() or "") for p in pdf.pages)
            if len(alt.strip()) > len(text.strip()):
                text = alt
        except Exception:
            pass
    return text


def _docx_text(data: bytes) -> str:
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise ApiError(422, "unreadable", "This DOCX could not be read.")
    parts = []
    for section in d.sections:  # contact details are often in headers/footers
        for part in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            try:
                parts.extend(p.text for p in part.paragraphs)
                for t in part.tables:
                    for row in t.rows:
                        parts.extend(c.text for c in row.cells)
            except Exception:
                pass
    parts.extend(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                parts.append(c.text)
    return "\n".join(parts)


def _doc_text(data: bytes) -> str:
    """Legacy .doc, best effort: pull long printable runs (ASCII and UTF-16LE)."""
    runs = re.findall(rb"[\x20-\x7e]{6,}", data)
    runs += [r.decode("utf-16le", "ignore").encode("utf-8") for r in re.findall(rb"(?:[\x20-\x7e]\x00){6,}", data)]
    return "\n".join(r.decode("utf-8", "ignore") for r in runs)


def extract_text(kind: str, data: bytes):
    """Returns (text, method). method 'gemini_vision' means text is empty and the PDF must be sent to Gemini."""
    if kind == "pdf":
        text = _normalise(_pdf_text(data))
        if len(text) < MIN_TEXT_CHARS:
            return "", "gemini_vision"
        return text[:MAX_TEXT_CHARS], "text"
    if kind == "docx":
        text = _normalise(_docx_text(data))
        if len(text) < 50:
            raise ApiError(422, "unreadable", "No readable text was found in this DOCX.")
        return text[:MAX_TEXT_CHARS], "text"
    text = _normalise(_doc_text(data))
    if len(text) < MIN_TEXT_CHARS:
        raise ApiError(422, "doc_unreadable",
                       "Legacy .doc could not be read. Please save it as PDF or DOCX and upload again.")
    return text[:MAX_TEXT_CHARS], "text"
